#!/usr/bin/env python3
"""Builds 'Getting Started with Claude and Obsidian.docx' using python-docx.

No Node/npm/pandoc/LibreOffice available in this environment, so this uses
python-docx directly instead of the docx-js approach. Numbered steps are
written as literal "1. " text (bold number) rather than Word auto-numbering,
so each section's list reliably restarts at 1 without needing custom
numbering.xml instances.
"""
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x5B, 0x3A, 0x8E)   # muted purple, restrained use only
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# ---- base page / font setup -------------------------------------------------
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(0.9)
section.bottom_margin = Inches(0.9)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after, bold in [
    ('Title', 28, ACCENT, 0, 4, True),
    ('Heading 1', 18, ACCENT, 20, 10, True),
    ('Heading 2', 14, INK, 14, 6, True),
    ('Heading 3', 12, INK, 10, 4, True),
]:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.page_break_before = False

doc.styles['List Bullet'].font.size = Pt(11)
doc.styles['List Bullet'].font.name = 'Calibri'

# ---- helpers -----------------------------------------------------------------

def set_cell_background(cell, hex_color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)

def remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)

def add_step(number, text_runs, note=None):
    """text_runs: list of (text, bold) tuples, or a plain string."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f'{number}.\t')
    r.bold = True
    r.font.color.rgb = ACCENT
    if isinstance(text_runs, str):
        text_runs = [(text_runs, False)]
    for txt, bold in text_runs:
        run = p.add_run(txt)
        run.bold = bold
    return p

def add_callout(title, lines, fill='EFEAF7'):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    remove_table_borders(table)
    set_cell_background(cell, fill)
    set_cell_margins(cell)
    cell.paragraphs[0].text = ''
    p0 = cell.paragraphs[0]
    if title:
        r0 = p0.add_run(title)
        r0.bold = True
        r0.font.color.rgb = ACCENT
        p0.paragraph_format.space_after = Pt(4)
        first_target = None
    else:
        first_target = p0  # reuse the empty first paragraph for the first content line
    for i, line in enumerate(lines):
        p = first_target if (first_target is not None and i == 0) else cell.add_paragraph(line)
        if p is first_target:
            p.text = line
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)  # breathing room after table

def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

def add_glossary_row(table, term, meaning, header=False):
    row = table.add_row()
    c1, c2 = row.cells
    c1.width = Inches(1.6)
    c2.width = Inches(4.9)
    set_cell_margins(c1, top=80, bottom=80)
    set_cell_margins(c2, top=80, bottom=80)
    if header:
        set_cell_background(c1, 'D9D2E9')
        set_cell_background(c2, 'D9D2E9')
        r1 = c1.paragraphs[0].add_run(term); r1.bold = True
        r2 = c2.paragraphs[0].add_run(meaning); r2.bold = True
    else:
        r1 = c1.paragraphs[0].add_run(term); r1.bold = True
        c2.paragraphs[0].add_run(meaning)

# ==============================================================================
# TITLE PAGE
# ==============================================================================
title = doc.add_paragraph(style='Title')
title.add_run('Getting Started with Claude and Obsidian')
title.paragraph_format.space_after = Pt(6)

sub = doc.add_paragraph()
sub.add_run('A beginner’s guide for Mac, Windows and Chromebook').font.size = Pt(13)
sub.runs[0].font.color.rgb = MUTED
sub.paragraph_format.space_after = Pt(18)

intro = doc.add_paragraph()
intro.add_run(
    'Obsidian is a free app for keeping notes as plain text files on your own computer. '
    'Claude is an AI assistant. Put them together and you get a notes system that an AI can '
    'actually read, search, and help you write in — not a chatbot you copy-paste into, but '
    'one that works alongside files you own and keep forever.'
)
intro.paragraph_format.space_after = Pt(4)

intro2 = doc.add_paragraph()
intro2.add_run(
    'This guide gets you from nothing installed to your first note, on a Mac, a Windows PC '
    'or a Chromebook. It takes about 15–20 minutes (about 30 on a Chromebook).'
)

add_callout('What you’ll need', [
    'A Mac, Windows PC or Chromebook with an internet connection. (A school or work Chromebook '
    'often cannot run Linux apps, which this guide needs.)',
    'About 15–20 minutes, uninterrupted (about 30 on a Chromebook).',
    'A paid Claude plan (Pro or higher). Claude’s Code feature, the part that works on your '
    'files, is not included in the free plan. Obsidian needs no account at all.',
    'Obsidian itself is free. The Claude plan is the only thing you pay for.',
])

# ==============================================================================
# WHAT ARE THESE TWO TOOLS
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.add_run('What you’re actually installing')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('Obsidian: where your notes live')
doc.add_paragraph(
    'Obsidian is a note-taking app that stores every note as a small text file on your own '
    'computer — not locked inside someone else’s server. A folder of notes is called a '
    '“vault.” Notes can link to each other by wrapping a title in double brackets, like '
    '[[this]], which is how Obsidian builds a web of connected ideas instead of a flat pile of '
    'documents.'
)

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('Claude: the assistant that can read and write those notes')
doc.add_paragraph(
    'Claude is Anthropic’s AI assistant. Used on its own, it’s a chat window. Used through '
    'its desktop app’s Code tab and pointed at a folder on your computer, it can read the '
    'files in that folder, write new ones, and edit existing ones — which is what turns an '
    'Obsidian vault from “notes you write” into “notes an assistant helps you keep.”'
)

add_callout('Why two separate apps, not one', [
    'Obsidian owns your notes: plain text files, yours forever, readable by any app, gone only '
    'if you delete them.',
    'Claude is the assistant that works on top of those files. If you ever stopped using Claude, '
    'every note would still open perfectly in Obsidian — nothing about your notes depends on it.',
], fill='F1F1F1')

# ==============================================================================
# PART 1: INSTALL OBSIDIAN
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.page_break_before = True
h.add_run('Part 1 — Install Obsidian')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On a Mac')
add_step(1, 'Open a web browser and go to obsidian.md.')
add_step(2, 'Click the purple “Download for macOS” button. The site detects your operating '
             'system automatically, so the button should already say macOS.')
add_step(3, 'Open the downloaded .dmg file from your Downloads folder, then drag the Obsidian '
             'icon into the Applications folder shown in that window.')
add_step(4, 'Open Obsidian from Applications (or Spotlight: press Cmd+Space, type Obsidian, '
             'press Enter). The first time, macOS may ask if you’re sure you want to open an '
             'app downloaded from the internet — click Open.')
add_callout('Don’t create a vault yet', [
    'Obsidian will offer to create or open a vault the moment it opens. Skip or cancel that for '
    'now — Part 3 walks through it deliberately, including where the vault should live.',
], fill='F1F1F1')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On Windows')
add_step(1, 'Open a web browser and go to obsidian.md.')
add_step(2, 'Click the “Download for Windows” button (the site detects Windows automatically).')
add_step(3, 'Open the downloaded installer from your Downloads folder and follow the prompts. '
             'It installs in a few seconds with no choices to make.')
add_step(4, 'Obsidian opens automatically when the install finishes. If Windows shows a '
             'SmartScreen warning about an unrecognized app, click “More info” then “Run '
             'anyway” — this is standard for new installers that haven’t built up download '
             'history yet, not a sign anything is wrong.')
add_callout('Don’t create a vault yet', [
    'Obsidian will offer to create or open a vault the moment it opens. Skip or cancel that for '
    'now — Part 3 walks through it deliberately, including where the vault should live.',
], fill='F1F1F1')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On a Chromebook')
doc.add_paragraph(
    'A Chromebook can run a small Linux computer inside it, and Obsidian runs there. You will '
    'paste a few short commands into a Terminal window; each one is written out for you.'
)
add_step(1, [('Turn on Linux: open ', False), ('Settings → About ChromeOS → Developers → Linux '
             'development environment → Set up', True), (' and follow the prompts (a few minutes). '
             'If you cannot find this setting, your Chromebook may be managed by a school or '
             'employer, and this guide will not work on it.', False)])
add_step(2, 'A Terminal window opens when Linux is ready (you can open it later by searching '
             '“Terminal” in the launcher). Paste this and press Enter: sudo apt update && sudo '
             'apt install -y curl git flatpak')
add_step(3, 'Paste: flatpak remote-add --if-not-exists flathub '
             'https://dl.flathub.org/repo/flathub.flatpakrepo')
add_step(4, 'Paste: flatpak install -y flathub md.obsidian.Obsidian')
add_step(5, 'Close the Terminal, shut Linux down (right-click Terminal in the launcher → Shut '
             'down Linux), reopen Terminal and start Obsidian with: flatpak run md.obsidian.Obsidian')
add_callout('If that does not work', [
    'The full Chromebook guide in the starter kit (chromebook-setup.md) has a second way to '
    'install Obsidian (an AppImage) and a fix for a common “libnss3” error. These Chromebook '
    'steps are drafted from the official instructions and are still being confirmed on a real '
    'Chromebook.',
], fill='F1F1F1')

# ==============================================================================
# PART 2: INSTALL CLAUDE
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.page_break_before = True
h.add_run('Part 2 — Install Claude')

doc.add_paragraph(
    'On a Mac or Windows PC you’ll use the Claude desktop app, which includes a “Code” tab — that’s the part '
    'that can open a folder on your computer and read or write files in it, rather than just chatting. '
    'On a Chromebook you’ll use the same thing from the Terminal.'
)

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On a Mac')
add_step(1, 'Go to claude.com/download.')
add_step(2, 'Click the macOS download button.')
add_step(3, 'Open the downloaded file and drag Claude into Applications, the same way you did '
             'for Obsidian.')
add_step(4, 'Open Claude from Applications, and sign in with your paid Claude account.')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On Windows')
add_step(1, 'Go to claude.com/download.')
add_step(2, 'Click the Windows download button.')
add_step(3, 'Run the installer and follow the prompts.')
add_step(4, 'Open Claude, and sign in with your paid Claude account.')

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On a Chromebook')
doc.add_paragraph(
    'On a Chromebook you use the command-line version of Claude, in the same Terminal window. '
    'It does the same job as the Code tab.'
)
add_step(1, 'In Terminal, paste this and press Enter (it shows no progress while it works): '
             'curl -fsSL https://claude.ai/install.sh | bash')
add_step(2, 'Close the Terminal, open it again, and type: claude --version. A version number '
             'means it worked.')
add_step(3, 'You sign in the first time you run it (Part 4). It opens a web page; if it does not, '
             'copy the link it prints into your Chrome browser.')

add_callout('A note on accounts', [
    'Obsidian needs no account at all — it just runs. Claude needs a paid plan (Pro or higher) '
    'to use its Code feature, because your conversations and settings are tied to your account.',
], fill='F1F1F1')

# ==============================================================================
# PART 3: CREATE YOUR FIRST VAULT
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.page_break_before = True
h.add_run('Part 3 — Create your first vault in Obsidian')

doc.add_paragraph(
    'A “vault” is just a folder. Obsidian doesn’t do anything special to it behind the scenes — '
    'every note inside is a plain .md (Markdown) text file you could open in any text editor. '
    'That’s deliberate: your notes are never trapped inside Obsidian.'
)

add_step(1, 'Open Obsidian. If it doesn’t prompt you automatically, click the vault icon in the '
             'very bottom-left corner and choose “Open another vault.”')
add_step(2, [('Choose ', False), ('Create new vault', True), ('.', False)])
add_step(3, [('Name it something like ', False), ('“My Notes”', True),
             (', and choose where it’s saved. The Desktop or Documents folder is a fine place '
              'to start — you can always move it later.', False)])
add_step(4, 'Click Create. You’ll land in an empty vault with a “Welcome” note already in it.')
add_callout('On a Chromebook', [
    'Choose the Linux home folder as the place to save the vault (it appears as “Linux files” '
    'in the Files app). Nothing syncs a Chromebook vault to the cloud for you: use ChromeOS '
    'Settings → About ChromeOS → Developers → Linux development environment → Backup & restore '
    'after big sessions, or see the full Chromebook guide for a Google Drive option.',
], fill='F1F1F1')
add_step(5, [('Create your first real note: click the ', False), ('new note', True),
             (' icon (top-left, looks like a page with a plus sign), then just start typing. '
              'Notes save automatically — there’s no save button.', False)])
add_step(6, [('Try a link: type ', False), ('[[', True), (' and Obsidian will suggest note '
              'names as you type — pick one (or type a new name and it creates that note for '
              'you). That double-bracket is the one piece of syntax worth learning first; '
              'everything else in Obsidian is just typing.', False)])

add_callout('The one habit worth building early', [
    'Link liberally. A note that links to three related notes is far more useful later than '
    'three unconnected notes — that web of links is most of what makes Obsidian better than a '
    'folder of Word documents.',
])

# ==============================================================================
# PART 4: CONNECT CLAUDE TO YOUR NOTES
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.page_break_before = True
h.add_run('Part 4 — Point Claude at your notes')

doc.add_paragraph(
    'This is the step that makes Claude more than a chat window: telling it which folder on '
    'your computer to work in.'
)

add_step(1, [('Open the Claude desktop app and click the ', False), ('Code', True),
             (' tab (it may also be labeled “Code” or show a folder-and-code icon, depending '
              'on your version).', False)])
add_step(2, 'Choose “Open folder” (or similar wording) and select your vault folder — the '
             'same one Obsidian created in Part 3.')
add_step(3, [('Type a message like ', False),
             ('“What files are in this folder?”', True),
             (' and send it. If Claude lists your notes back to you, it’s working.', False)])

h2 = doc.add_paragraph(style='Heading 2'); h2.add_run('On a Chromebook')
add_step(1, 'In Terminal, go to your vault folder. If you named it “My Notes” and saved it in '
             'Linux files, type: cd ~/"My Notes"')
add_step(2, 'Type: claude and press Enter. Sign in when it asks (it opens a web page).')
add_step(3, [('Then type a message like ', False), ('“What files are in this folder?”', True),
             (' If Claude lists your notes, it’s working. Keep this Terminal window open next '
              'to Obsidian; changes show up in both.', False)])

add_callout('What this actually gives you', [
    'Claude can now read your notes for context, create new ones, and edit existing ones when '
    'you ask it to — in plain English, the same way you’d ask a person. It never changes a file '
    'without being asked to, and anything it writes is a normal text file you can open, edit, '
    'or delete in Obsidian exactly like one you wrote yourself.',
], fill='F1F1F1')

# ==============================================================================
# PART 5: TRY IT
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.add_run('Part 5 — A first thing to try')

doc.add_paragraph('With Claude still open on your vault folder, try asking it something like:')

add_callout('', [
    '“Read the note I just created and suggest two things I could link it to.”',
    'or',
    '“Create a note called Reading List with a few books I mention, one per line.”',
])

doc.add_paragraph(
    'Then go check Obsidian — the change will already be there. That back-and-forth, editing '
    'the same files from two directions, is the whole idea.'
)

# ==============================================================================
# PART 6: GOING FURTHER
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.add_run('Part 6 — Want to go further?')

doc.add_paragraph(
    'Once the basics feel comfortable, some people set up a more structured vault: separate '
    'space for their own freeform notes versus a hierarchy an AI assistant maintains for '
    'ongoing projects, plus a house style file so Claude follows the same conventions every '
    'session instead of being told from scratch each time. That’s a deliberate system, not '
    'something to build on day one — but if it’s of interest later, ask whoever shared this '
    'guide with you about the “Second Brain” setup, or ask Claude directly to help you design '
    'one once you’re past the basics here.'
)

# ==============================================================================
# GLOSSARY
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.page_break_before = True
h.add_run('Quick glossary')

table = doc.add_table(rows=0, cols=2)
table.autofit = False
add_glossary_row(table, 'Term', 'What it means', header=True)
add_glossary_row(table, 'Vault', 'A folder of notes that Obsidian treats as one connected '
                  'collection. Just a normal folder underneath.')
add_glossary_row(table, 'Note', 'A single Markdown (.md) text file. One idea per note is the '
                  'usual rule of thumb.')
add_glossary_row(table, 'Markdown', 'A simple way of formatting plain text (like **bold** or '
                  '# heading) that any text editor can read, not just Obsidian.')
add_glossary_row(table, 'Wikilink', 'Double brackets around a note title, like [[Note Name]], '
                  'that creates a clickable link to that note.')
add_glossary_row(table, 'Terminal / Linux', 'On a Chromebook, a text window where you paste '
                  'commands to the small Linux computer inside it. Obsidian and Claude run there.')
add_glossary_row(table, 'Code tab', 'The part of the Claude desktop app that can open a folder '
                  'and read or write files in it, instead of just chatting.')

doc.add_paragraph()  # spacer

# ==============================================================================
# TROUBLESHOOTING
# ==============================================================================
h = doc.add_paragraph(style='Heading 1')
h.add_run('If something doesn’t work')

add_bullets([
    'Obsidian won’t open on Windows and shows a security warning: click “More info,” then “Run '
    'anyway.” This is standard for a newly downloaded installer.',
    'Mac says the app is from an “unidentified developer”: right-click the app in Applications, '
    'choose Open, then confirm — this is only needed the first time.',
    'Claude’s Code tab can’t find your notes: double-check you opened the same folder Obsidian '
    'created, not a different one, and that you selected the folder itself rather than a file '
    'inside it.',
    'Chromebook: there is no “Linux development environment” setting: your Chromebook is '
    'probably managed by a school or work account, which can switch Linux off.',
    'A note you expected to see in Obsidian isn’t there after Claude created it: click into the '
    'file explorer pane on the left and refresh, or close and reopen the vault — Obsidian '
    'usually picks up new files immediately, but occasionally needs a nudge.',
])

doc.save('Getting Started with Claude and Obsidian.docx')
print('done')

